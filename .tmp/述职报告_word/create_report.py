from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('信息技术部负责人履职情况报告（初稿）.docx')

# Palette / named design override:
# standard_business_brief + cn_official_a4 (A4, Chinese official-document typography)
NAVY = RGBColor(31, 78, 121)
DARK = RGBColor(31, 31, 31)
MUTED = RGBColor(100, 100, 100)
LIGHT = 'EEF3F8'
BORDER = 'CBD5E1'
RED = RGBColor(156, 0, 6)

FONT_BODY = 'Songti SC'
FONT_HEADING = 'STHeiti'
FONT_LATIN = 'Calibri'


def set_east_asia_font(run, font_name, latin_name=None):
    latin_name = latin_name or font_name
    run.font.name = latin_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), latin_name)
    rFonts.set(qn('w:hAnsi'), latin_name)
    rFonts.set(qn('w:cs'), latin_name)


def set_style_font(style, east_asia, size, bold=False, color=None, latin=FONT_LATIN):
    style.font.name = latin
    style._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia)
    style._element.rPr.rFonts.set(qn('w:ascii'), latin)
    style._element.rPr.rFonts.set(qn('w:hAnsi'), latin)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = color


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_table_borders(table, color=BORDER, size='6'):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = borders.find(qn(f'w:{edge}'))
        if e is None:
            e = OxmlElement(f'w:{edge}')
            borders.append(e)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), size)
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)


def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(dxa))
    tcW.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths)))
    tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.find(qn('w:tblInd'))
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(w)); grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])


def set_keep_with_next(p, keep=True):
    p.paragraph_format.keep_with_next = keep


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])
    set_east_asia_font(run, FONT_BODY)
    run.font.size = Pt(9)


def add_total_pages_field(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' NUMPAGES '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])
    set_east_asia_font(run, FONT_BODY)
    run.font.size = Pt(9)


def add_body(doc, text, bold_prefix=None, indent=True, after=6, keep=False):
    p = doc.add_paragraph(style='正文')
    p.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True; set_east_asia_font(r1, FONT_BODY)
        r2 = p.add_run(text[len(bold_prefix):]); set_east_asia_font(r2, FONT_BODY)
    else:
        r = p.add_run(text); set_east_asia_font(r, FONT_BODY)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(style='一级标题')
    p.add_run(text)
    for r in p.runs: set_east_asia_font(r, FONT_HEADING)
    set_keep_with_next(p)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style='二级标题')
    p.add_run(text)
    for r in p.runs: set_east_asia_font(r, FONT_HEADING)
    set_keep_with_next(p)
    return p


def add_h3(doc, text):
    p = doc.add_paragraph(style='三级标题')
    p.add_run(text)
    for r in p.runs: set_east_asia_font(r, FONT_BODY)
    set_keep_with_next(p)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [8500])
    set_table_borders(table, color='B8C7D9', size='8')
    cell = table.cell(0,0); shade_cell(cell, LIGHT); set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.2
    r = p.add_run(label + '：'); r.bold = True; r.font.color.rgb = NAVY; set_east_asia_font(r, FONT_HEADING)
    r2 = p.add_run(text); set_east_asia_font(r2, FONT_BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text):
    p = doc.add_paragraph(style='项目符号')
    r = p.add_run(text); set_east_asia_font(r, FONT_BODY)
    return p


def add_signature(doc):
    p = doc.add_paragraph(style='正文')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('述职人：【姓名待确认】'); set_east_asia_font(r, FONT_BODY); r.bold = True
    p2 = doc.add_paragraph(style='正文')
    p2.paragraph_format.space_after = Pt(0)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p2.add_run('【日期待确认】'); set_east_asia_font(r, FONT_BODY)


doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210); sec.page_height = Mm(297)
sec.top_margin = Mm(25.4); sec.bottom_margin = Mm(25.4)
sec.left_margin = Mm(27); sec.right_margin = Mm(27)
sec.header_distance = Mm(12.5); sec.footer_distance = Mm(12.5)

# Defaults and styles
normal = doc.styles['Normal']
set_style_font(normal, FONT_BODY, 12, color=DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.4
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

styles = doc.styles
for name in ['正文','一级标题','二级标题','三级标题','项目符号','校核正文']:
    if name not in styles:
        styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

body = styles['正文']; set_style_font(body, FONT_BODY, 12, color=DARK)
body.paragraph_format.line_spacing = 1.4
body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
body.paragraph_format.widow_control = True

h1 = styles['一级标题']; set_style_font(h1, FONT_HEADING, 16, bold=True, color=NAVY)
h1.paragraph_format.space_before = Pt(16); h1.paragraph_format.space_after = Pt(8)
h1.paragraph_format.line_spacing = 1.2; h1.paragraph_format.keep_with_next = True; h1.paragraph_format.widow_control = True

h2 = styles['二级标题']; set_style_font(h2, FONT_HEADING, 13, bold=True, color=NAVY)
h2.paragraph_format.space_before = Pt(12); h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.line_spacing = 1.25; h2.paragraph_format.keep_with_next = True; h2.paragraph_format.widow_control = True

h3 = styles['三级标题']; set_style_font(h3, FONT_BODY, 12, bold=True, color=DARK)
h3.paragraph_format.space_before = Pt(8); h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.line_spacing = 1.25; h3.paragraph_format.keep_with_next = True

bullet = styles['项目符号']; set_style_font(bullet, FONT_BODY, 11.5, color=DARK)
bullet.paragraph_format.left_indent = Mm(10); bullet.paragraph_format.first_line_indent = Mm(-5)
bullet.paragraph_format.tab_stops.add_tab_stop(Mm(10))
bullet.paragraph_format.space_after = Pt(4); bullet.paragraph_format.line_spacing = 1.3

check = styles['校核正文']; set_style_font(check, FONT_BODY, 10.5, color=DARK)
check.paragraph_format.line_spacing = 1.25; check.paragraph_format.space_after = Pt(4)

# Header/footer
header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = hp.add_run('信息技术部负责人履职情况报告｜内部拟稿')
set_east_asia_font(r, FONT_HEADING); r.font.size = Pt(9); r.font.color.rgb = MUTED
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('第 '); set_east_asia_font(r, FONT_BODY); r.font.size = Pt(9); r.font.color.rgb = MUTED
add_page_field(fp)
r = fp.add_run(' 页 / 共 '); set_east_asia_font(r, FONT_BODY); r.font.size = Pt(9); r.font.color.rgb = MUTED
add_total_pages_field(fp)
r = fp.add_run(' 页'); set_east_asia_font(r, FONT_BODY); r.font.size = Pt(9); r.font.color.rgb = MUTED

# Opening block: memo masthead adapted for formal Chinese report
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(22); p.paragraph_format.space_after = Pt(6)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('信息技术部负责人履职情况报告')
set_east_asia_font(r, FONT_HEADING); r.font.size = Pt(24); r.bold = True; r.font.color.rgb = DARK
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(18)
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run('（【述职周期待确认】）')
set_east_asia_font(r, FONT_BODY); r.font.size = Pt(12); r.font.color.rgb = MUTED

# Metadata, real table because repeated label/value fields
meta = doc.add_table(rows=3, cols=2)
set_table_geometry(meta, [1550, 6950])
set_table_borders(meta, color='D5DEE8', size='5')
meta_data = [('报告人','【姓名待确认】'),('职务','信息技术部副部长（主持工作）【待确认】'),('汇报对象','【主送对象待确认】')]
for i,(label,value) in enumerate(meta_data):
    set_cant_split(meta.rows[i])
    c0,c1 = meta.rows[i].cells
    shade_cell(c0, LIGHT)
    for c in (c0,c1):
        set_cell_margins(c, 100, 140, 100, 140); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c0.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    rr=p.add_run(label); set_east_asia_font(rr,FONT_HEADING); rr.bold=True; rr.font.size=Pt(10.5); rr.font.color.rgb=NAVY
    p=c1.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    rr=p.add_run(value); set_east_asia_font(rr,FONT_BODY); rr.font.size=Pt(10.5)

doc.add_paragraph().paragraph_format.space_after = Pt(0)
add_callout(doc, '拟稿口径', '依据《信息技术部工作情况汇报》（2026年3月）整理，按部门负责人个人述职口径表述；未核实的信息均以【待确认】标识。')

add_body(doc, '【主送对象待确认】：', indent=False, after=8)
add_body(doc, '按照公司统一安排，现将本人【述职周期待确认】履职情况报告如下。')

add_h1(doc, '一、岗位履职总体情况')
add_body(doc, '【述职周期待确认】，在公司党委和经营管理层的领导下，本人立足信息技术部负责人岗位，紧紧围绕公司建设“一流司库型财务公司”的总体目标，落实集团“数字航发”战略，坚持以战略规划为牵引、以项目建设为抓手、以数据治理为主线、以安全稳定为底线，统筹推进信息科技治理、数字化转型、重大项目建设、网络与数据安全、系统运行维护、外包服务管理及人才队伍建设等工作。')
add_body(doc, '履职过程中，本人重点抓好信息化规划和年度计划落实，统筹信息化项目全生命周期管理，推进信息技术架构优化和国产化适配，完善网络安全、数据治理、外包管理等制度机制，协调信息科技预算和供应商管理，组织开展信息科技风险防控和业务连续性建设，推动信息技术由传统的业务支撑工具逐步向业务赋能和价值创造平台转变。')
add_body(doc, '在部门人员和专业资源相对有限的情况下，本人注重发挥内部员工的管理主导作用和外包人员的专业支撑作用，推动形成分工协作、相互补位的工作机制，较好完成了公司交办的各项信息科技工作任务。')

add_h1(doc, '二、主要履职成效')
add_h2(doc, '（一）强化战略统筹，信息科技治理体系进一步完善')
add_body(doc, '一是持续完善信息科技顶层设计。组织编制《“十五五”信息科技专项规划》，形成“1248”发展思路；推动完成数字化转型战略规划与蓝图设计，明确智慧金融中枢、穿透式风控、产融协同、自主可控、数字文化五大愿景，构建业务、应用、数据、人工智能、技术和安全“六位一体”的转型架构，为后续数字化建设提供总体指引。')
add_body(doc, '二是持续健全信息科技制度体系。围绕数据治理、网络安全、项目建设、运行维护和外包管理等领域，推动信息科技管理制度修订完善。“十四五”期间累计修订相关制度32项，信息科技管理的制度化、规范化水平得到提升。')
add_body(doc, '三是不断优化治理机制。推动公司成立数字化转型工作领导小组，引入专业机构开展数字化现状评估和实施路径规划，组织开展数据治理专题培训。推动自研信息科技管理系统，建立9大类、30余个数字化管理台账，进一步增强信息科技工作的可视化、规范化和过程管控能力。')

add_h2(doc, '（二）加强项目统筹，重大信息化建设取得积极成效')
add_body(doc, '一是推进司库体系建设。组织推动覆盖集团全级次的司库管理平台建设，支撑集团资金集中管理。平台获集团验收评级优秀，资金集中度和可归集资金集中度均达到95%。2025年进一步推进平台升级，支持金融资源整合和跨境监测，提升集团资金管理的数字化、集约化水平。')
add_body(doc, '二是推进核心系统国产化替代。组织完成司库商密网国产化建设，推动国产芯片、麒麟操作系统、国产数据库等软硬件应用，建成60平方米B级机房；推进司库管理平台及数据仓库40亿条数据由Oracle数据库向达梦数据库迁移；组织新一代核心业务系统硬件集成和上线试运行，核心系统信创替代比例达到60%。相关信创建设经验入选集团典型案例。')
add_body(doc, '三是推进穿透式监管能力建设。组织建设集团穿透式监管系统一期，覆盖国资委“11+4”核心风险模型，实现资金流向穿透展示和相关监管模型线上运行。根据后续业务和监管需求，启动二期项目，拟进一步建设薪资监控、智能风险识别、异常交易监测、系统集成和产融协同等功能。')
add_body(doc, '四是持续完善金融服务平台。完成国产化新一代票据系统建设，支撑票据承兑、转贴现和再贴现等业务开展；推进监管报送系统升级改造，实现246张报表及1104、EAST等10余类监管数据自动化报送；完成国产桌面虚拟化替代，办公终端国产化覆盖率达到100%。')

add_h2(doc, '（三）深化数据治理，数据基础能力持续增强')
add_body(doc, '一是推进企业级数据平台建设。构建“司库中台+监管报送仓”数据底座，完成3大类、29个主题、424个标准项梳理，形成161个关键指标。数据平台承载数据量由2023年的8亿条增长至2024年的30亿条，核心业务数据入湖率达到80%。2025年发布包含6大类、200余项内容的数据资产目录，数据资源管理和服务能力进一步增强。')
add_body(doc, '二是推动数据贯通共享。利用单向网闸技术，打通公司商密网至集团涉密网的数据通道，累计导入司库数据5000万条；推进司库管理驾驶舱与集团数据中心贯通，数据流转效率提升50%以上；建设集团级单向网闸通道，为公司与集团ERP、供应链等系统开展安全数据交换提供支撑。')
add_body(doc, '三是加强数据标准和质量管理。推动发布《司库信息分类与属性》数据标准，并将其融入集团AEOS体系；组织编制数据分类分级目录，划分L1—L4级主题域72个，核心数据定级完成率达到90%；持续推进数据标准化治理，核心数据标准化覆盖率达到85%，监管报表自动化率提升至90%。')
add_body(doc, '四是积极推动自主创新。组织自研国资委司库监管检核工具，内置970余项检核规则，实现监管数据报送零差错，节约采购成本100万元，在提高工作效率的同时增强了团队自主研发能力。')

add_h2(doc, '（四）守牢安全底线，网络和数据安全保障能力不断提升')
add_body(doc, '一是压实网络安全责任。持续完善网络安全和数据安全管理体系，组织开展安全风险评估、漏洞排查整改、应急演练和安全培训。“十四五”期间未发生重大网络安全事件，核心系统持续开展等级保护三级测评和商用密码应用安全性评估，2024年密评结果为“良好”。')
add_body(doc, '二是抓好重大网络安全保障。组织参加集团HW行动，2023年连续防守15天，封堵攻击IP一万余个；2024年引入“金钟罩”主动防御系统，完成30天连续防守，封堵攻击IP五万余个，实现零失陷，公司成为集团9家满分单位之一。连续两年参加科工局“军工网盾”超限演练，响应及时率达到100%。')
add_body(doc, '三是强化业务连续性建设。完成司库商密网、金融服务平台、OA等6个核心系统等保三级测评及复测，推动档案、投资等系统通过等保二级测评；建成同城A级容灾备份中心，开展“两地三中心”容灾体系实战化验证，核心系统恢复时间目标达到RTO不超过30分钟、RPO接近零。')
add_body(doc, '四是细化数据安全管理。部署数据库加密脱敏设备，推动测试数据“可用不可见”；建立数据全生命周期安全管理机制，常态化开展数据安全风险评估，持续落实数据分类分级和差异化保护措施。')

add_h2(doc, '（五）夯实运维基础，系统服务保障质效稳步提升')
add_body(doc, '一是保障核心系统稳定运行。2023—2025年，核心业务系统、司库系统等关键应用年可用率保持在99.99%，故障处理及时率达到100%。组织开展系统巡检、数据备份、升级优化和故障处置，2023年完成系统巡检近500次、备份数据库文件1.8万套次，2024年完成系统升级优化94次、处理终端问题700余次。')
add_body(doc, '二是提升用户服务水平。建立7×24小时线上支持机制，三年累计通过电话、微信等渠道解决成员单位问题3000余次，客户满意度保持在98%以上。推动完成司库管理驾驶舱67项优化，进一步提升系统使用体验和决策支持能力。')
add_body(doc, '三是做好重大时期保障。组织完成历年“两会”、年终结算、人行ACS切换等重要时期和重大事项保障任务；支撑集团完成财政部专线统一接入；协助部署境外安全保障应急指挥系统，为公司及集团相关业务稳定运行提供技术支持。')

add_h2(doc, '（六）探索技术创新，数字化赋能取得阶段性进展')
add_body(doc, '一是推进人工智能应用试点。围绕智能风控、资产配置、制度问答等场景开展人工智能应用探索。基于知识图谱的智能风控助手获得集团数字化转型创新二等奖，2025年推动人工智能应用在11类场景试点落地。')
add_body(doc, '二是推动流程自动化应用。推进RPA“数字员工”在银企对账、监管报表生成等高频场景应用，实现部分业务7×24小时自动化作业，提升工作效率，降低人工操作差错风险。')
add_body(doc, '三是开展技术架构前瞻布局。在数字化转型蓝图中规划“云原生+人工智能中台+数据湖”技术架构，组织开展核心系统云原生改造预研，为后续系统智能化升级打好基础。')
add_body(doc, '四是加强团队能力建设。通过项目实践、外部培训和内部技术交流等方式，着力培养既懂金融业务又掌握数字技术的复合型人才。2023—2025年，团队累计参加外部培训50余场，专业能力和协作水平持续提升。')

add_h1(doc, '三、履职中存在的主要不足')
add_body(doc, '在总结成绩的同时，本人也清醒认识到，对照公司数字化转型要求和信息科技管理目标，工作中仍存在以下不足。')
add_h2(doc, '（一）数字化转型统筹的深度仍需加强')
add_body(doc, '前期工作更多聚焦核心业务系统建设，经营管理流程尚未得到全面、系统梳理。预算管理、绩效考核、内控审批、决策支持等管理流程与业务流程之间仍存在衔接不够紧密的问题，对标集团AEOS 2.0“全流程、全要素”管理要求还有差距。')
add_body(doc, '司库平台与集团ERP、供应链、合同等系统尚未实现全链路数据贯通，业财资税融合仍有断点。部分风险管理功能主要集中于事后分析，事前预警和事中干预能力仍需提高。作为部门负责人，本人在跨部门统筹、流程协同和系统性推动方面还需进一步加大力度。')
add_h2(doc, '（二）数据治理和价值挖掘仍需深化')
add_body(doc, '部分边缘系统数据尚未完全纳入统一治理范围，跨平台数据标准还未实现全域统一，数据孤岛问题尚未彻底解决。现有数据应用仍以统计查询和报表展示为主，面向流动性预测、产业链风险传导和经营决策的智能分析能力相对不足，数据资产的业务价值尚未得到充分释放。')
add_body(doc, '同时，随着数据应用范围扩大，数据分类分级、权限控制、安全共享和全生命周期管理面临更高要求，数据治理与数据安全之间的协同机制还需进一步完善。')
add_h2(doc, '（三）人工智能应用转化能力仍显不足')
add_body(doc, '现有人工智能应用大多处于单点试点阶段，在信贷审批、流动性预测、智能客服等核心业务场景尚未形成规模化应用。从技术验证到业务流程嵌入，再到价值评价和推广复制，仍缺少成熟、清晰的转化路径。国产化环境下的系统性能调优和智能应用组件适配也需要持续投入。')
add_h2(doc, '（四）部门资源配置与承担职能不完全匹配')
add_body(doc, '信息技术部同时承担系统建设、运行维护、网络安全、基础设施、数据治理、人工智能应用、流程梳理及外包管理等多项职责，工作领域多、专业跨度大。部门在数据科学、人工智能算法等领域的专业人才储备不足，部分岗位对外部技术力量的依赖程度较高，存在知识沉淀和知识转移不足的风险。')
add_body(doc, '本人在内部人才培养、关键岗位备份和外包成果转化方面虽然开展了一些工作，但还没有形成足够系统的机制，部门专业化、精细化管理能力仍需提升。')

add_h1(doc, '四、下一步履职思路和重点工作')
add_body(doc, '2026年是“十五五”开局之年。下一步，本人将围绕公司数字化转型总体部署，坚持问题导向和目标导向，重点做好以下工作。')
add_h2(doc, '（一）抓好规划落实和科技治理')
add_body(doc, '按照“十五五”信息科技专项规划和数字化转型蓝图，进一步细化年度任务、责任分工和实施计划，加强重大项目统筹和全过程管理。完善项目立项、需求评审、开发测试、上线验收及后评价机制，提高信息科技投资效益。')
add_body(doc, '围绕AEOS 2.0建设要求，配合相关部门推进经营管理流程梳理，促进预算、绩效、内控、决策支持等管理流程与业务流程衔接，推动流程优化成果向信息系统固化。')
add_h2(doc, '（二）推进重点项目建设')
add_body(doc, '加快推进集团穿透式监管系统二期建设，做好采购、开发、系统联调和上线试运行等工作，按照现有计划推动项目于四季度上线。')
add_body(doc, '全力保障新一代核心业务系统平稳切换上线，加强上线前测试、数据迁移、用户培训、应急预案和运行监控，确保新旧系统平稳过渡。持续推进核心系统和周边系统国产化适配及性能优化，提高关键技术自主可控水平。')
add_body(doc, '按照现有项目计划推进新OA系统建设，做好需求确认、系统开发、测试上线和用户培训，推动办公流程线上化、移动化，并逐步加强与财务、人力等系统的集成。')
add_h2(doc, '（三）深化数据治理和数据安全管理')
add_body(doc, '组织开展数据治理专项攻坚，重点推进信贷、资金、票据等核心领域的数据标准统一和数据质量提升。完善数据资源目录、数据质量问题清单和整改责任机制，推动更多业务数据纳入统一治理体系。')
add_body(doc, '牵头做好集团财务管理域数据整理相关工作，推动财务数据资源目录编制、数据口径统一和数据整合技术方案制定，配合打通与集团相关系统的数据接口。')
add_body(doc, '落实数据安全管理要求，开展数据安全全面自查和风险评估，完善数据分类分级台账、核心数据及重要数据清单，加强敏感数据脱敏、加密、访问控制和安全监测，守住不发生重大数据安全事件的底线。')
add_h2(doc, '（四）稳妥推进人工智能应用')
add_body(doc, '围绕资金预测、信贷风控和流动性管理等重点场景，推进财务垂域模型项目论证和高质量数据集建设，做好数据清洗、标注、质量校验和应用验证。')
add_body(doc, '坚持业务需求牵引，加强人工智能项目的安全性、适用性和实际效益评估，推动成熟试点逐步从单点验证转向业务流程应用，避免脱离实际需求的重复建设。')
add_h2(doc, '（五）提升安全运维保障能力')
add_body(doc, '持续做好网络设备巡检、漏洞扫描、补丁更新和安全策略优化，组织开展红蓝对抗和应急演练，完善网络安全应急响应机制，做好HW行动备战和重要时期网络安全保障。')
add_body(doc, '加强系统运行监测、数据备份和恢复测试，持续检验业务连续性和容灾保障能力。完善外包服务商准入、尽职调查、过程监督、绩效评价和退出机制，强化外包人员权限管理、操作审计及知识转移，降低供应链和外包服务风险。')
add_h2(doc, '（六）加强队伍建设和能力培养')
add_body(doc, '结合部门职责和项目建设需要，加强数据治理、人工智能、网络安全、项目管理等方面专业能力建设。通过项目实战、专业培训、岗位交流和内部分享，培养金融业务与数字技术相结合的复合型人才。')
add_body(doc, '进一步明确内部人员和外包人员职责边界，推动关键技术、系统文档和运维经验向内部团队沉淀，逐步提升自主建设、自主运维和自主风险处置能力。')

add_h1(doc, '五、需进一步研究和协调的事项')
add_body(doc, '一是进一步研究数据管理组织机制。随着数据治理、数据资产管理和人工智能应用任务不断增加，现有组织架构与专业化管理要求之间存在一定差距，可结合公司整体机构设置和职责分工，对数据管理职能的组织模式开展专题研究。')
add_body(doc, '二是加强复合型人才保障。根据数字化转型实际需要，研究信息科技、数据科学和人工智能等专业人才的引进与培养机制，完善业务与技术人员双向交流和岗位培养安排。')
add_body(doc, '三是加强数字化转型资金统筹。根据“十五五”期间相关工程建设安排，统筹考虑年度项目预算的连续性和重点项目的资金保障，降低因预算衔接影响项目进度的风险。')
add_body(doc, '四是加强集团层面的数据共享和系统对接协调。公司与集团ERP、供应链、合同等系统开展数据贯通，需要进一步明确接口标准、数据口径、共享权限和安全责任。')

add_body(doc, '以上是本人的履职情况。下一步，本人将进一步增强责任意识，认真履行岗位职责，持续提升信息科技治理和数字化转型工作质效，为公司建设“一流司库型财务公司”提供安全、稳定、高效的科技支撑。')
add_signature(doc)

# Separate reviewer-only appendix
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run('拟稿校核清单')
set_east_asia_font(r, FONT_HEADING); r.font.size=Pt(18); r.bold=True; r.font.color.rgb=NAVY
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(14)
r=p.add_run('（报送或印发前删除本页）'); set_east_asia_font(r,FONT_BODY); r.font.size=Pt(10.5); r.font.color.rgb=RED; r.bold=True

items = [
('述职主体','现有材料显示常允为信息技术部副部长并主持工作，但尚不能确认本报告是否以常允个人名义述职。'),
('述职周期','建议明确为“2025年度”“2023—2025年”或“十四五期间”，避免年度履职与跨年度成果口径混用。'),
('人员数量','原材料称“正式员工3人、外包人员3人，共6人”，但人员表列出4名内部人员和3名外包人员，共7人，需核准。'),
('政策依据','“2026年监管重点要求”“国资委最新报送要求”等依据未列明正式文件名称，使用前应据实核验。'),
('评价与荣誉','“行业率先”“获监管部门充分肯定”“集团9家满分单位之一”“入选集团典型案例”等表述，建议核对验收文件或通报。'),
('年度计划数据','原材料中“XX家、XX个、XX类、XX%、XX项”等数据尚未确定，正式报告中应补充真实数据或删除。'),
('技术指标','“RPO接近零”建议核对技术检测或演练报告中的正式表述。'),
('统计口径','“三年累计修订制度32项”等统计口径需确认是否包括新建、修订和废止。'),
('其他履职内容','如单位述职要求包含党建、廉洁从业、意识形态或“一岗双责”等内容，需另行提供经核验材料。'),
]
t = doc.add_table(rows=1, cols=3)
set_table_geometry(t, [650, 1650, 6200]); set_table_borders(t)
h = t.rows[0]; set_repeat_table_header(h); set_cant_split(h)
for i,txt in enumerate(['序号','待核事项','核验说明']):
    c=h.cells[i]; shade_cell(c,'DDE7F1'); set_cell_margins(c,120,120,120,120); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(txt); set_east_asia_font(r,FONT_HEADING); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=NAVY
for idx,(name,desc) in enumerate(items,1):
    row=t.add_row(); set_cant_split(row)
    vals=[str(idx),name,desc]
    for j,val in enumerate(vals):
        c=row.cells[j]; set_cell_margins(c,100,120,100,120); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.2
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j<2 else WD_ALIGN_PARAGRAPH.LEFT
        r=p.add_run(val); set_east_asia_font(r,FONT_BODY); r.font.size=Pt(10)

p=doc.add_paragraph(style='校核正文')
p.paragraph_format.space_before=Pt(12)
r=p.add_run('使用提示：'); set_east_asia_font(r,FONT_HEADING); r.bold=True; r.font.color.rgb=NAVY
r=p.add_run('本文件仅供拟稿和校对辅助。最终文种、事实数据、权限和单位格式，应由述职人及业务负责人审核确认；本助手不代替签发、盖章、发布、报送或提交。'); set_east_asia_font(r,FONT_BODY)

# Document metadata and compatibility options
props=doc.core_properties
props.title='信息技术部负责人履职情况报告（初稿）'
props.subject='依据2026年3月信息技术部工作情况汇报整理'
props.author='公文写作助手（拟稿）'
props.keywords='述职报告, 信息技术部, 初稿, 待核'

settings = doc.settings._element
compat = settings.find(qn('w:compat'))
if compat is None:
    compat = OxmlElement('w:compat'); settings.append(compat)
doNotExpand = OxmlElement('w:doNotExpandShiftReturn'); compat.append(doNotExpand)

# Ensure all table rows can split only when explicitly desired (all here kept intact)
for table in doc.tables:
    for row in table.rows:
        set_cant_split(row)

# Last orphan control
for p in doc.paragraphs:
    p.paragraph_format.widow_control = True

# Save
if OUT.exists(): OUT.unlink()
doc.save(OUT)
print(OUT.resolve())
